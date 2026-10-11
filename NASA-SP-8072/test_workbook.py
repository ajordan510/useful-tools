"""Read browser-generated XLSX files independently and verify every exported value."""
import json, math, zipfile
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parent
checks=[]
def close(actual,expected):
    assert isinstance(actual,(int,float)), (actual,expected)
    assert math.isclose(actual,expected,rel_tol=1e-13,abs_tol=1e-12), (actual,expected)
for name in ['results','modern','deflected','literal']:
    path=root/'_qa'/f'{name}.xlsx'
    expected=json.loads(path.with_suffix('.json').read_text(encoding='utf-8'))
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        assert not any(b'<x:f' in z.read(n) for n in z.namelist() if n.startswith('xl/worksheets/'))
    wb=load_workbook(path,data_only=False)
    assert wb.sheetnames==['Summary','Bands','Inputs']
    summary,bands,inputs=(wb[s] for s in wb.sheetnames)
    for sheet in wb:
        assert sheet.freeze_panes=='A8'
        assert sheet.auto_filter.ref is not None
        assert sheet.sheet_view.showGridLines is False
        assert sheet['A2'].value==expected['c']['name']
        assert not any(c.data_type in ('f','e') for row in sheet for c in row)
    nband=len(expected['bands']);nr=len(expected['receivers'])
    assert summary.max_row==7+nr
    assert bands.max_row==7+nr*nband
    for i,p in enumerate(expected['receivers']):
        row=list(summary.iter_rows(min_row=8+i,max_row=8+i,values_only=True))[0]
        assert row[0]==p['name']
        for v,key in zip(row[1:],['x','y','oa','la','design','rms','peak','minR','minTheta','maxTheta']):close(v,p[key])
        for j,b in enumerate(expected['bands']):
            row=list(bands.iter_rows(min_row=8+i*nband+j,max_row=8+i*nband+j,values_only=True))[0]
            assert row[0]==p['name']
            for a,v in zip(row[1:7],[p['x'],p['y'],b['f'],b['lo'],b['hi'],p['levels'][j]]):close(a,v)
            close(row[7],10*math.log10(expected['bandPower'][j]/1e-12))
            f2=b['f']**2
            close(row[8],20*math.log10(12194**2*f2*f2/((f2+20.6**2)*math.sqrt((f2+107.7**2)*(f2+737.9**2))*(f2+12194**2)))+2)
            close(row[9],p['levels'][j]+expected['c']['margin'])
            close(row[10],p['oa']);close(row[11],p['la']);close(row[12],p['rms'])
    records={row[0]:row[1] for row in inputs.iter_rows(min_row=8,values_only=True)}
    for k,v in expected['c'].items():
        if k!='receivers':assert records[k]==(v if v!='' else None), (k,records[k],v)
    assert json.loads(records['Case JSON'])==expected['c']
    close(records['Derived acoustic power'],expected['W'])
    close(records['Derived acoustic power level'],expected['Lw'])
    close(records['Effective diameter'],expected['D'])
    close(records['Core length'],expected['core'])
    close(records['Band power fraction'],expected['capture'])
    assert records['Band count']==nband and records['Receiver count']==nr
    assert [records[f'Model warning {i+1}'] for i in range(len(expected['warnings']))]==expected['warnings']
    assert 'Reference 1' in records
    checks.append({'case':name,'receiver_count':nr,'band_rows':nr*nband,'passed':True})
    print(f'PASS {name}: {nr} receiver summaries and {nr*nband} band rows; typed values, inputs, warnings and literal text')
(root/'reference/workbook-test-results.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
